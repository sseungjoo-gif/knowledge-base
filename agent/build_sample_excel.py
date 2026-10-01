"""agent/sample_data_gen.py의 데이터를 다운로드용 엑셀 파일로 내보낸다.

실행:
    cd "D:\\knowledge\\files (2)"
    py -m agent.build_sample_excel
"""

from openpyxl import Workbook
from openpyxl.styles import Font

from . import sample_data_gen as gen


def _autosize(ws):
    for col in ws.columns:
        length = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(60, length + 2)


def _header(ws, headers):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)


def build_workbook(out_path: str):
    data = gen.build_all()
    wb = Workbook()

    # ── 안내 ──
    ws = wb.active
    ws.title = "안내"
    rows = [
        ["워커힐 호텔 가상 샘플 데이터 (2023-01 ~ 2026-12)"],
        [],
        ["※ 전부 가상(synthetic) 데이터입니다. 실제 워커힐 실적이 아닙니다."],
        [],
        ["적용한 규칙(배부 기준 등)"],
        ["항목", "값", "설명"],
        ["가용객실", "500실 × 해당월 일수", ""],
        ["OTA 채널수수료율", "18%", "OTA 판매분 매출에만 적용"],
        ["객실 부문 직접비", "OTA수수료 + 객실매출 × 14.5%", "14.5%는 하우스키핑·소모품 등 고정 운영비 가정"],
        ["F&B 부문 직접비", "F&B매출 × 70%", ""],
        ["연회(Banquet) 부문 직접비", "연회매출 × 62.5%", ""],
        ["기타 부문 직접비", "기타매출 × 25%", ""],
        ["공통비 배부 기준", "부문 매출 비례", "관리 8.82% · 마케팅 5.88% · 시설 4.41% · 에너지 5.88% (부문 매출 합계 대비)"],
        ["표준 식자재 원가율", "32%", "전 기간 고정 정책값. 실제원가율은 월별로 변동"],
        ["고객 통합키 매핑율", "70% (가정)", "실측치가 아니라 샘플 생성을 위한 가정값"],
        [],
        ["2026-12는 1.1(매출·수익성 지표 흐름) 문서의 예시 숫자와 정확히 일치하도록 고정했고,"],
        ["나머지 47개월은 위 규칙 + 추세(OCC 68%→80%, ADR 25.5만→30만, OTA비중 35%→58%) + 계절성 + 약간의 무작위 변동으로 생성했습니다."],
        [],
        ["시트 구성"],
        ["객실_월별", "월별 가용객실·판매객실·매출·OTA판매"],
        ["부문손익_월별", "부문(객실·F&B·연회·기타)별 매출·직접비"],
        ["미배부공통비_월별", "관리·마케팅·시설·에너지"],
        ["FB원가_월별", "F&B 식자재 실제원가·표준원가율"],
        ["시장벤치마크_월별", "가상 Comp Set OCC·ADR·RevPAR"],
        ["고객마스터", "가상 고객 300명 (세그먼트·국적·등급·통합키매핑여부)"],
        ["고객연간이용", "고객별 연도별 숙박일수·객실매출·F&B매출·채널"],
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A5"].font = Font(bold=True)
    ws["A19"].font = Font(bold=True)
    _autosize(ws)

    # ── 객실_월별 ──
    ws = wb.create_sheet("객실_월별")
    _header(ws, ["년월", "일수", "가용객실", "판매객실", "객실매출", "OTA판매객실", "OTA매출", "OCC", "ADR"])
    for lbl, r in data["room_inventory"].items():
        avail = gen.ROOM_COUNT * r["days"]
        ws.append([lbl, r["days"], avail, r["sold_rooms"], r["revenue"],
                   r["ota_sold_rooms"], r["ota_revenue"],
                   round(r["sold_rooms"] / avail, 4), round(r["revenue"] / r["sold_rooms"])])
    _autosize(ws)

    # ── 부문손익_월별 ──
    ws = wb.create_sheet("부문손익_월별")
    _header(ws, ["년월", "부문", "매출", "직접비", "부문이익(배부전)"])
    for lbl, depts in data["department_pl"].items():
        for dept, v in depts.items():
            ws.append([lbl, dept, v["revenue"], v["direct_cost"], v["revenue"] - v["direct_cost"]])
    _autosize(ws)

    # ── 미배부공통비_월별 ──
    ws = wb.create_sheet("미배부공통비_월별")
    _header(ws, ["년월", "관리", "마케팅", "시설", "에너지", "합계"])
    for lbl, u in data["undistributed_expense"].items():
        ws.append([lbl, u["관리"], u["마케팅"], u["시설"], u["에너지"], sum(u.values())])
    _autosize(ws)

    # ── FB원가_월별 ──
    ws = wb.create_sheet("FB원가_월별")
    _header(ws, ["년월", "F&B매출", "실제식자재원가", "실제원가율", "표준원가율", "괴리(%p)"])
    for lbl, f in data["fb_cost"].items():
        actual_rate = f["actual_food_cost"] / f["food_revenue"]
        ws.append([lbl, f["food_revenue"], f["actual_food_cost"], round(actual_rate, 4),
                   f["standard_food_cost_rate"], round((actual_rate - f["standard_food_cost_rate"]) * 100, 2)])
    _autosize(ws)

    # ── 시장벤치마크_월별 ──
    ws = wb.create_sheet("시장벤치마크_월별")
    _header(ws, ["년월", "시장OCC", "시장ADR", "시장RevPAR"])
    for lbl, m in data["market_benchmark"].items():
        ws.append([lbl, m["market_occ"], m["market_adr"], m["market_revpar"]])
    _autosize(ws)

    # ── 고객마스터 ──
    ws = wb.create_sheet("고객마스터")
    _header(ws, ["고객ID", "세그먼트", "국적", "회원등급", "통합키매핑여부"])
    for c in data["customers"]:
        ws.append([c["customer_id"], c["segment"], c["nationality"], c["membership_tier"], c["integration_key_matched"]])
    _autosize(ws)

    # ── 고객연간이용 ──
    ws = wb.create_sheet("고객연간이용")
    _header(ws, ["고객ID", "연도", "숙박일수", "객실매출", "F&B매출", "채널"])
    for u in data["customer_usage"]:
        ws.append([u["customer_id"], u["year"], u["room_nights"], u["room_revenue"], u["fb_revenue"], u["channel"]])
    _autosize(ws)

    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    path = build_workbook("워커힐_샘플데이터_2023-2026.xlsx")
    print(f"저장됨: {path}")
